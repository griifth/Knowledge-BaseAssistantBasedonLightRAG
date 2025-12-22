"""
Chat-related routes for LightRAG API.
Provides file parsing and web search for chat context injection.
"""

import asyncio
import os
import httpx
from io import BytesIO
from typing import Optional, List
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel, Field
from lightrag.api.utils_api import get_combined_auth_dependency
from lightrag.utils import logger


router = APIRouter(tags=["chat"])

# 博查搜索 API 配置
BOCHA_API_URL = "https://api.bocha.cn/v1/web-search"


class ParsedFileResponse(BaseModel):
    """Response model for parsed file content"""
    filename: str = Field(description="Original filename")
    content_type: str = Field(description="MIME type of the file")
    size: int = Field(description="File size in bytes")
    content: str = Field(description="Extracted text content")
    error: Optional[str] = Field(default=None, description="Error message if parsing failed")


# 搜索相关数据模型
class WebSearchRequest(BaseModel):
    """Request model for web search"""
    query: str = Field(description="Search query")
    count: int = Field(default=8, ge=1, le=50, description="Number of results to return")
    freshness: str = Field(default="noLimit", description="Time range: noLimit, oneDay, oneWeek, oneMonth, oneYear")
    summary: bool = Field(default=True, description="Whether to include summary in results")


class WebSearchResult(BaseModel):
    """Single search result"""
    title: str = Field(description="Page title")
    url: str = Field(description="Page URL")
    snippet: str = Field(description="Short description")
    summary: Optional[str] = Field(default=None, description="Full summary if requested")
    siteName: Optional[str] = Field(default=None, description="Site name")
    datePublished: Optional[str] = Field(default=None, description="Publish date")


class WebSearchResponse(BaseModel):
    """Response model for web search"""
    query: str = Field(description="Original search query")
    results: List[WebSearchResult] = Field(description="Search results")
    totalMatches: int = Field(default=0, description="Total estimated matches")
    error: Optional[str] = Field(default=None, description="Error message if search failed")


def _extract_pdf(file_bytes: bytes, password: str = None) -> str:
    """Extract PDF content using pypdf."""
    from pypdf import PdfReader

    pdf_file = BytesIO(file_bytes)
    reader = PdfReader(pdf_file)

    if reader.is_encrypted:
        if not password:
            raise Exception("PDF is encrypted but no password provided")
        decrypt_result = reader.decrypt(password)
        if decrypt_result == 0:
            raise Exception("Incorrect PDF password")

    content = ""
    for page in reader.pages:
        text = page.extract_text()
        if text:
            content += text + "\n"

    return content.strip()


def _extract_docx(file_bytes: bytes) -> str:
    """Extract DOCX content including tables."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    docx_file = BytesIO(file_bytes)
    doc = Document(docx_file)

    content_parts = []

    for element in doc.element.body:
        if element.tag.endswith("p"):
            paragraph = Paragraph(element, doc)
            text = paragraph.text
            if text.strip():
                content_parts.append(text)

        elif element.tag.endswith("tbl"):
            table = Table(element, doc)
            for row in table.rows:
                row_text = [cell.text for cell in row.cells]
                if any(cell.strip() for cell in row_text):
                    content_parts.append("\t".join(row_text))

    return "\n".join(content_parts)


def _extract_pptx(file_bytes: bytes) -> str:
    """Extract PPTX content."""
    from pptx import Presentation

    pptx_file = BytesIO(file_bytes)
    prs = Presentation(pptx_file)
    
    content_parts = []
    for slide_num, slide in enumerate(prs.slides, 1):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text)
        if slide_text:
            content_parts.append(f"[Slide {slide_num}]\n" + "\n".join(slide_text))

    return "\n\n".join(content_parts)


def _extract_xlsx(file_bytes: bytes) -> str:
    """Extract XLSX content."""
    from openpyxl import load_workbook

    xlsx_file = BytesIO(file_bytes)
    wb = load_workbook(xlsx_file)

    content_parts = []
    for sheet in wb.worksheets:
        sheet_content = [f"[Sheet: {sheet.title}]"]
        for row in sheet.iter_rows(values_only=True):
            row_text = [str(cell) if cell is not None else "" for cell in row]
            if any(cell.strip() for cell in row_text):
                sheet_content.append("\t".join(row_text))
        if len(sheet_content) > 1:
            content_parts.append("\n".join(sheet_content))

    return "\n\n".join(content_parts)


def _extract_text(file_bytes: bytes) -> str:
    """Extract plain text content."""
    try:
        return file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        try:
            return file_bytes.decode("gbk")
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1")


# Supported file extensions and their extractors
EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".pptx": _extract_pptx,
    ".xlsx": _extract_xlsx,
    ".txt": _extract_text,
    ".md": _extract_text,
    ".csv": _extract_text,
    ".json": _extract_text,
    ".xml": _extract_text,
    ".html": _extract_text,
    ".htm": _extract_text,
    ".log": _extract_text,
    ".yaml": _extract_text,
    ".yml": _extract_text,
}


def create_chat_routes(api_key: Optional[str] = None):
    """Create chat-related routes."""
    
    combined_auth = get_combined_auth_dependency(api_key)

    @router.post(
        "/chat/parse-file",
        response_model=ParsedFileResponse,
        dependencies=[Depends(combined_auth)],
        summary="Parse file for chat context",
        description="Upload and parse a file to extract text content for chat context injection. "
                    "Supports PDF, Word, PowerPoint, Excel, and various text formats."
    )
    async def parse_file_for_chat(
        file: UploadFile = File(..., description="File to parse")
    ):
        """
        Parse an uploaded file and extract its text content.
        
        Supported formats:
        - PDF (.pdf)
        - Word (.docx)
        - PowerPoint (.pptx)
        - Excel (.xlsx)
        - Text files (.txt, .md, .csv, .json, .xml, .html, .log, .yaml, .yml)
        
        Returns the extracted text content that can be injected into chat context.
        """
        filename = file.filename or "unknown"
        content_type = file.content_type or "application/octet-stream"
        
        # Get file extension
        ext = ""
        if "." in filename:
            ext = "." + filename.rsplit(".", 1)[-1].lower()
        
        # Check if supported
        if ext not in EXTRACTORS:
            return ParsedFileResponse(
                filename=filename,
                content_type=content_type,
                size=0,
                content="",
                error=f"Unsupported file type: {ext}. Supported: {', '.join(EXTRACTORS.keys())}"
            )
        
        try:
            # Read file content
            file_bytes = await file.read()
            file_size = len(file_bytes)
            
            # Size limit: 20MB
            if file_size > 20 * 1024 * 1024:
                return ParsedFileResponse(
                    filename=filename,
                    content_type=content_type,
                    size=file_size,
                    content="",
                    error="File too large. Maximum size is 20MB."
                )
            
            # Extract content using appropriate extractor
            extractor = EXTRACTORS[ext]
            content = await asyncio.to_thread(extractor, file_bytes)
            
            # Truncate if too long (limit to 50000 chars for context)
            max_chars = 50000
            if len(content) > max_chars:
                content = content[:max_chars] + f"\n\n... [Content truncated, showing first {max_chars} characters]"
            
            logger.info(f"Parsed file {filename}: {len(content)} chars extracted")
            
            return ParsedFileResponse(
                filename=filename,
                content_type=content_type,
                size=file_size,
                content=content
            )
            
        except Exception as e:
            logger.error(f"Error parsing file {filename}: {str(e)}")
            return ParsedFileResponse(
                filename=filename,
                content_type=content_type,
                size=0,
                content="",
                error=f"Failed to parse file: {str(e)}"
            )

    @router.post(
        "/chat/web-search",
        response_model=WebSearchResponse,
        dependencies=[Depends(combined_auth)],
        summary="Web search for chat context",
        description="Search the web using Bocha AI API and return results for chat context injection."
    )
    async def web_search_for_chat(request: WebSearchRequest):
        """
        Search the web using Bocha AI API.
        
        Requires BOCHA_API_KEY to be set in environment variables.
        
        Returns search results that can be injected into chat context.
        """
        # 从环境变量获取 API Key
        bocha_api_key = os.getenv("BOCHA_API_KEY")
        
        if not bocha_api_key:
            return WebSearchResponse(
                query=request.query,
                results=[],
                error="Web search is not configured. Please set BOCHA_API_KEY in environment variables."
            )
        
        try:
            async with httpx.AsyncClient(timeout=30.0) as client:
                response = await client.post(
                    BOCHA_API_URL,
                    headers={
                        "Authorization": f"Bearer {bocha_api_key}",
                        "Content-Type": "application/json"
                    },
                    json={
                        "query": request.query,
                        "count": request.count,
                        "freshness": request.freshness,
                        "summary": request.summary
                    }
                )
                
                if response.status_code != 200:
                    error_msg = f"Search API returned {response.status_code}"
                    try:
                        error_data = response.json()
                        if "message" in error_data:
                            error_msg = error_data["message"]
                    except:
                        pass
                    
                    logger.error(f"Bocha search error: {error_msg}")
                    return WebSearchResponse(
                        query=request.query,
                        results=[],
                        error=error_msg
                    )
                
                data = response.json()
                
                # 解析响应
                results = []
                total_matches = 0
                
                if data.get("code") == 200 and data.get("data"):
                    search_data = data["data"]
                    web_pages = search_data.get("webPages", {})
                    total_matches = web_pages.get("totalEstimatedMatches", 0)
                    
                    for item in web_pages.get("value", []):
                        results.append(WebSearchResult(
                            title=item.get("name", ""),
                            url=item.get("url", ""),
                            snippet=item.get("snippet", ""),
                            summary=item.get("summary"),
                            siteName=item.get("siteName"),
                            datePublished=item.get("datePublished") or item.get("dateLastCrawled")
                        ))
                
                logger.info(f"Web search for '{request.query}': {len(results)} results")
                
                return WebSearchResponse(
                    query=request.query,
                    results=results,
                    totalMatches=total_matches
                )
                
        except httpx.TimeoutException:
            logger.error(f"Web search timeout for query: {request.query}")
            return WebSearchResponse(
                query=request.query,
                results=[],
                error="Search request timed out"
            )
        except Exception as e:
            logger.error(f"Web search error: {str(e)}")
            return WebSearchResponse(
                query=request.query,
                results=[],
                error=f"Search failed: {str(e)}"
            )
    
    return router

